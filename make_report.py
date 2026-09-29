# -*- coding: utf-8 -*-
"""
把实验结果整理成 Word 实验报告（.docx）。

以表格、配图和必要注释为主体，不做大段文字陈述。

用法：
    python run.py            # 先跑实验
    python make_report.py    # 生成 成果汇总.docx

依赖：python-docx   （pip install python-docx）
"""
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'output')
DST = os.path.join(HERE, '成果汇总.docx')


def set_font(run, ascii_font='Times New Roman', cn_font='宋体', size=None,
             bold=None, color=None):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para(doc, text='', size=10.5, bold=False, cn='宋体', align=None, color=None,
         space_after=4, indent=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    if text:
        set_font(p.add_run(text), 'Times New Roman', cn, size, bold, color)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), 'Times New Roman', '黑体', 12.5, True)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(text), 'Times New Roman', '宋体', 9.5)
    return p


def image_row(doc, items, total_width=6.25, cap=None):
    items = [(f, c) for f, c in items if os.path.exists(os.path.join(OUT, f))]
    if not items:
        return
    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    w = total_width / len(items) - 0.06
    for cell, (fn, c) in zip(t.rows[0].cells, items):
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(1)
        cp.add_run().add_picture(os.path.join(OUT, fn), width=Inches(w))
    if cap:
        caption(doc, cap)


def add_image(doc, fn, width_in, cap=None):
    p = os.path.join(OUT, fn)
    if not os.path.exists(p):
        return
    doc.add_picture(p, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
    if cap:
        caption(doc, cap)


def no_split(table):
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def make_table(doc, headers, rows, widths=None, bold_last=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
        set_font(c.paragraphs[0].add_run(h), 'Times New Roman', '黑体', 9, True)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            bold = bold_last and ri == len(rows) - 1
            set_font(p.add_run(str(v)), 'Times New Roman', '宋体', 9, bold)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    no_split(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def main():
    p = os.path.join(OUT, '结果数据.json')
    if not os.path.exists(p):
        print('没找到 output/结果数据.json，请先运行 python run.py')
        return
    D = json.load(open(p, encoding='utf-8'))
    P, A, B = D['params'], D['route_a'], D['route_b']
    comp, hist = D['compare'], D['convergence']

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    base = doc.styles['Normal']
    base.font.name = 'Times New Roman'
    base.font.size = Pt(10.5)
    base.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    try:
        tg = doc.styles['Table Grid']
        tg.font.name = 'Times New Roman'
        tg.font.size = Pt(9)
        tg.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    except Exception:
        pass

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run('户型自动生成：两条技术路线的对比'), 'Times New Roman', '黑体', 17, True)
    para(doc, '实验报告', size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    para(doc, '路线 A：Power Diagram（权重反解）　|　路线 B：墙图 + 递归切分（约束搜索）',
         size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    # ---- 一、方法 ----
    heading(doc, '一、方法')
    make_table(doc, ['', '路线 A：Power Diagram', '路线 B：墙图 + 递归切分'],
               [['表示', '加权 Voronoi 胞腔', '墙图：顶点=墙交点，边=墙段，面=房间'],
                ['生成', '站点位置给定，反解权重控制面积',
                 '按面积比例递归切分矩形'],
                ['面积控制', '反馈迭代反解，误差约 1e-4',
                 '按比例切分，误差为 0'],
                ['约束', '未实现', '搜索房间顺序，满足邻接约束'],
                ['墙体', '区域边界，方向任意', '全部横平竖直']],
               widths=[0.85, 2.55, 2.8])

    # ---- 二、路线 A 结果 ----
    heading(doc, '二、路线 A：Power Diagram 生成结果')
    image_row(doc, [('A_户型_一室一厅.png', ''), ('A_户型_两室一厅.png', '')],
              cap='（a）一室一厅　63.0 m²　　（b）两室一厅　93.5 m²')
    add_image(doc, 'A_户型_三室一厅.png', 5.4, '（c）三室一厅　117.0 m²')
    make_table(doc, ['户型', '房间数', '面积占比总误差', '外接矩形填充率'],
               [[r['name'], len(r['rooms']), '%.6f' % r['err'], '%.3f' % r['fill']]
                for r in A],
               widths=[1.5, 1.3, 2.0, 1.5])
    para(doc, '外接矩形填充率 = 区域实际面积 ÷ 其外接矩形面积。矩形为 1.0，'
              '倾斜多边形小于 1。路线 A 只有 0.73～0.80，说明房间是斜边的多边形。',
         indent=True, size=10, space_after=8)

    # ---- 三、权重的作用 ----
    heading(doc, '三、路线 A：权重的作用')
    add_image(doc, 'A_对比_Voronoi与Power.png', 6.2,
              '（d）等权重（Voronoi）：面积由几何位置被动决定　|　'
              '（e）优化权重：面积精确达标')
    make_table(doc, ['房间', '目标占比', '等权重（Voronoi）', '优化权重'],
               [[r['name'], '%.2f%%' % (r['target'] * 100),
                 '%.2f%%' % (r['equal'] * 100), '%.2f%%' % (r['optimized'] * 100)]
                for r in comp['rooms']])
    para(doc, '面积占比总误差：等权重 %.4f → 优化权重 %.6f（改善约 %d 倍）'
         % (comp['err_equal'], comp['err_optimized'],
            round(comp['err_equal'] / max(comp['err_optimized'], 1e-9))),
         size=10, bold=True, space_after=8)

    # ---- 四、收敛 ----
    heading(doc, '四、路线 A：权重优化的收敛过程')
    add_image(doc, 'A_对比_收敛曲线.png', 5.4,
              '（f）面积误差随迭代次数下降（%s）' % comp['preset'])
    para(doc, '起始误差 %.4f，迭代 %d 次后降至 %.6f。全程单调下降、无超调，'
              '前 20 次即完成主要收敛。' % (hist[0], len(hist) - 1, hist[-1]),
         indent=True, space_after=8)

    # ---- 五、路线 B 结果 ----
    heading(doc, '五、路线 B：墙图路线生成结果')
    image_row(doc, [('B_墙图_一室一厅.png', ''), ('B_墙图_两室一厅.png', '')],
              cap='（g）一室一厅　63.0 m²　　（h）两室一厅　93.5 m²')
    add_image(doc, 'B_墙图_三室一厅.png', 5.4, '（i）三室一厅　117.0 m²')
    para(doc, '路线 B 的房间全部为矩形、墙体横平竖直，外墙上绘制了窗户符号，'
              '门开在相邻房间的公共墙上。', indent=True, size=10, space_after=6)
    make_table(doc, ['户型', '面积误差', '约束违反', '搜索次数', '耗时 (s)',
                     '墙图规模'],
               [[r['name'], '%.6f' % r['err'], r['violations'], r['tries'],
                 '%.2f' % r['sec'], '%d顶点/%d边' % (r['n_vertices'], r['n_edges'])]
                for r in B],
               widths=[1.1, 1.3, 0.95, 0.95, 0.9, 1.35])

    # ---- 六、对比 ----
    heading(doc, '六、两条路线的定量对比')
    add_image(doc, '对比_两条路线.png', 6.3,
              '（j）路线 A（左）：不规则多边形　|　路线 B（右）：矩形房间')
    a, b = A[1], B[1]
    make_table(doc, ['指标', '路线 A（Power Diagram）', '路线 B（墙图）'],
               [['面积占比总误差', '%.6f' % a['err'], '%.6f' % b['err']],
                ['外接矩形填充率', '%.3f' % a['fill'], '1.000'],
                ['邻接约束', '未支持', '已满足（0 违反）'],
                ['计算耗时', '约 1.1 秒（300 次迭代）', '%.2f 秒（%d 次搜索）'
                 % (b['sec'], b['tries'])],
                ['房间形状', '不规则多边形，边数不定', '全部为矩形']],
               widths=[1.5, 2.35, 2.35])
    para(doc, '路线 B 在形状规整度、面积精度、约束支持三方面都更接近真实户型；'
              '路线 A 的优势在于几何表达更灵活，可用于非矩形场地或斜向分区。',
         indent=True, space_after=8)

    # ---- 七、问题记录 ----
    heading(doc, '七、问题记录')
    make_table(doc, ['#', '现象', '原因', '修法'],
               [['1', '门洞开到了外墙上',
                 'np.roll 是环绕操作，膨胀时把图像另一侧的区域卷了过来，'
                 '外墙边缘凭空出现"邻居"',
                 '膨胀前先 padding；并限制门洞只能开在距外墙一定距离的内墙上'],
                ['2', '生成一张户型要 14 秒',
                 '权重优化迭代 300 次，每次都跑全分辨率',
                 '优化改在低分辨率（220 px）上进行，收敛后按全分辨率重算，'
                 '耗时降至 1.1 秒'],
                ['3', '收敛曲线剧烈震荡',
                 '反馈增益过大，学习率 1.2 时误差冲到 0.88，超调约 7 倍',
                 '学习率降至 0.4，全程单调下降'],
                ['4', '邻接约束一开始全部违反',
                 '递归切分的结果依赖房间排列顺序，而顺序决定了谁和谁相邻',
                 '对房间顺序做随机搜索，以"约束违反数 + 形状惩罚"为评分，'
                 '3～6 次即找到满足全部约束的解']],
               widths=[0.3, 1.3, 2.55, 2.05])

    doc.save(DST)
    print('已生成：%s   (%.0f KB)' % (DST, os.path.getsize(DST) / 1024))


if __name__ == '__main__':
    main()
